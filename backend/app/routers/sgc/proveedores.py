"""
SGC — Gestión de Proveedores
============================
SGC crea y administra el catálogo de proveedores.
OC Automatizaciones lee de aquí para poblar el selector de proveedores.

Extracción automática
---------------------
El endpoint POST /extraer acepta PDF, Excel o Word y pre-llena los campos
del formulario usando regex. Los patrones están en _parsear_campos_rut() y
se calibrarán cuando llegue el formato oficial de evaluación de proveedores.
"""

from __future__ import annotations

import io
import re
import uuid
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from pydantic import BaseModel
from sqlmodel import Session, select

from app.core.deps import get_current_user
from app.models.sgc import ProveedorSGC
from app.models.user import User
from app.sgc_database import get_sgc_db

router = APIRouter(prefix="/proveedores", tags=["SGC - Proveedores"])


# ── Schemas ───────────────────────────────────────────────────────────────────

class ProveedorSGCCreate(BaseModel):
    nombre: str
    nit: Optional[str] = None
    representante_legal: Optional[str] = None
    email: Optional[str] = None
    telefono: Optional[str] = None
    direccion: Optional[str] = None
    ciudad: Optional[str] = None
    categoria: Optional[str] = None
    observaciones: Optional[str] = None


class ProveedorSGCUpdate(BaseModel):
    nombre: Optional[str] = None
    nit: Optional[str] = None
    representante_legal: Optional[str] = None
    email: Optional[str] = None
    telefono: Optional[str] = None
    direccion: Optional[str] = None
    ciudad: Optional[str] = None
    categoria: Optional[str] = None
    observaciones: Optional[str] = None
    activo: Optional[bool] = None


class ProveedorSGCRead(BaseModel):
    id: uuid.UUID
    nombre: str
    nit: Optional[str]
    representante_legal: Optional[str]
    email: Optional[str]
    telefono: Optional[str]
    direccion: Optional[str]
    ciudad: Optional[str]
    categoria: Optional[str]
    activo: bool
    observaciones: Optional[str]
    creado_por_id: Optional[int]
    created_at: datetime
    updated_at: Optional[datetime]

    class Config:
        from_attributes = True


class ExtraccionRUTResult(BaseModel):
    """Resultado de la extracción automática desde documento del proveedor."""
    nombre: Optional[str] = None
    nit: Optional[str] = None
    representante_legal: Optional[str] = None
    email: Optional[str] = None
    telefono: Optional[str] = None
    direccion: Optional[str] = None
    ciudad: Optional[str] = None
    nombre_archivo: str = ""
    campos_encontrados: int = 0


# ── Helpers de extracción de texto ────────────────────────────────────────────
# Reutiliza la misma lógica que OC cotizaciones para leer PDF/Excel/Word.

def _extraer_texto(contenido: bytes, ext: str) -> str:
    if ext == "pdf":
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(contenido)) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        except Exception:
            return ""
    if ext in ("xlsx", "xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts = [str(c) for c in row if c is not None]
                    if parts:
                        lines.append("  ".join(parts))
            return "\n".join(lines)
        except Exception:
            return ""
    if ext == "docx":
        try:
            from docx import Document
            doc = Document(io.BytesIO(contenido))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception:
            return ""
    return ""


def _parsear_campos_rut(texto: str) -> ExtraccionRUTResult:
    """
    Extrae campos de un documento de proveedor (RUT u otro formato).

    ⚠️  Esta función se calibrará cuando llegue el formato oficial de
    evaluación de proveedores. Los patrones actuales cubren el RUT
    colombiano estándar y documentos de presentación comunes.

    Para agregar un nuevo patrón, añadirlo a la lista correspondiente.
    """

    def find(patterns: list[str]) -> Optional[str]:
        for pat in patterns:
            m = re.search(pat, texto, re.IGNORECASE | re.MULTILINE)
            if m:
                val = m.group(1).strip()
                if val:
                    return val[:300]
        return None

    # ── NIT ───────────────────────────────────────────────────────────────────
    nit = find([
        r"N\.?I\.?T\.?[:\s#]*(\d[\d.\-]+[-]\d)",
        r"NIT[:\s#]*(\d{3}[.\s]?\d{3}[.\s]?\d{3}[-]?\d)",
    ])

    # ── Razón social / Nombre empresa ─────────────────────────────────────────
    nombre = find([
        r"RAZ[OÓ]N\s+SOCIAL[:\s]+(.{3,200}?)(?:\n|$)",
        r"NOMBRE(?:S)?\s+(?:Y\s+APELLIDOS\s+O\s+RAZ[OÓ]N\s+SOCIAL|COMPLETO|S)?[:\s]+(.{3,200}?)(?:\n|$)",
        r"ENTIDAD[:\s]+(.{3,200}?)(?:\n|$)",
    ])

    # ── Representante legal ───────────────────────────────────────────────────
    representante = find([
        r"REPRESENTANTE\s+LEGAL[:\s]+(.{3,200}?)(?:\n|$)",
        r"APODERADO[:\s]+(.{3,200}?)(?:\n|$)",
        r"GERENTE[:\s]+(.{3,200}?)(?:\n|$)",
    ])

    # ── Email ─────────────────────────────────────────────────────────────────
    email = find([
        r"(?:E-?MAIL|CORREO\s+ELECTR[OÓ]NICO)[:\s]+([\w.\-+]+@[\w\-]+\.[\w.\-]+)",
        r"([\w.\-+]+@[\w\-]+\.[\w.\-]+)",
    ])

    # ── Teléfono ──────────────────────────────────────────────────────────────
    telefono = find([
        r"(?:TEL[EÉ]FONO|CELULAR|M[OÓ]VIL|PBX|TEL\.?)[:\s]+([\d\s\-+()]{7,20})",
        r"(?:TEL|CEL)[:\s]+([\d\s\-+()]{7,20})",
    ])

    # ── Dirección ─────────────────────────────────────────────────────────────
    direccion = find([
        r"DIRECCI[OÓ]N[:\s]+(.{5,200}?)(?:\n|$)",
        r"DOMICILIO[:\s]+(.{5,200}?)(?:\n|$)",
        r"DIRECCI[OÓ]N\s+SECCIONAL[:\s]+(.{5,200}?)(?:\n|$)",
    ])

    # ── Ciudad / Municipio ────────────────────────────────────────────────────
    ciudad = find([
        r"(?:CIUDAD|MUNICIPIO)[:\s]+(.{3,100}?)(?:\n|$)",
        r"(?:CIUDAD|MUNICIPIO)\s*/\s*DEPARTAMENTO[:\s]+(.{3,100}?)(?:\n|$)",
    ])

    campos = sum(
        1 for v in [nit, nombre, representante, email, telefono, direccion, ciudad]
        if v is not None
    )

    return ExtraccionRUTResult(
        nombre=nombre,
        nit=nit,
        representante_legal=representante,
        email=email,
        telefono=telefono,
        direccion=direccion,
        ciudad=ciudad,
        campos_encontrados=campos,
    )


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.get("", response_model=list[ProveedorSGCRead])
def list_proveedores(
    solo_activos: bool = Query(default=False),
    current_user: User = Depends(get_current_user),
    sgc_db: Session = Depends(get_sgc_db),
):
    """Lista proveedores. SGC ve todos; por defecto muestra activos e inactivos."""
    query = select(ProveedorSGC)
    if solo_activos:
        query = query.where(ProveedorSGC.activo == True)  # noqa: E712
    query = query.order_by(ProveedorSGC.nombre)
    return sgc_db.exec(query).all()


@router.get("/{proveedor_id}", response_model=ProveedorSGCRead)
def get_proveedor(
    proveedor_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    sgc_db: Session = Depends(get_sgc_db),
):
    proveedor = sgc_db.get(ProveedorSGC, proveedor_id)
    if not proveedor:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Proveedor no encontrado.")
    return proveedor


@router.post("", response_model=ProveedorSGCRead, status_code=status.HTTP_201_CREATED)
def create_proveedor(
    payload: ProveedorSGCCreate,
    current_user: User = Depends(get_current_user),
    sgc_db: Session = Depends(get_sgc_db),
):
    existing = sgc_db.exec(
        select(ProveedorSGC).where(ProveedorSGC.nombre == payload.nombre)
    ).first()
    if existing:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"Ya existe un proveedor con el nombre '{payload.nombre}'.",
        )
    proveedor = ProveedorSGC(
        **payload.model_dump(),
        creado_por_id=current_user.id,
    )
    sgc_db.add(proveedor)
    sgc_db.commit()
    sgc_db.refresh(proveedor)
    return proveedor


@router.put("/{proveedor_id}", response_model=ProveedorSGCRead)
def update_proveedor(
    proveedor_id: uuid.UUID,
    payload: ProveedorSGCUpdate,
    current_user: User = Depends(get_current_user),
    sgc_db: Session = Depends(get_sgc_db),
):
    proveedor = sgc_db.get(ProveedorSGC, proveedor_id)
    if not proveedor:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Proveedor no encontrado.")

    data = payload.model_dump(exclude_unset=True)

    if "nombre" in data and data["nombre"] != proveedor.nombre:
        conflict = sgc_db.exec(
            select(ProveedorSGC).where(ProveedorSGC.nombre == data["nombre"])
        ).first()
        if conflict:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=f"Ya existe un proveedor con el nombre '{data['nombre']}'.",
            )

    for field, value in data.items():
        setattr(proveedor, field, value)
    proveedor.updated_at = datetime.now(timezone.utc)

    sgc_db.add(proveedor)
    sgc_db.commit()
    sgc_db.refresh(proveedor)
    return proveedor


@router.patch("/{proveedor_id}/toggle-activo", response_model=ProveedorSGCRead)
def toggle_activo(
    proveedor_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    sgc_db: Session = Depends(get_sgc_db),
):
    """Activa o desactiva un proveedor. Los inactivos no aparecen en OC."""
    proveedor = sgc_db.get(ProveedorSGC, proveedor_id)
    if not proveedor:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Proveedor no encontrado.")

    proveedor.activo = not proveedor.activo
    proveedor.updated_at = datetime.now(timezone.utc)
    sgc_db.add(proveedor)
    sgc_db.commit()
    sgc_db.refresh(proveedor)
    return proveedor


@router.post("/extraer", response_model=ExtraccionRUTResult)
async def extraer_documento(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
):
    """
    Extrae campos del proveedor desde un PDF, Excel o Word.
    No guarda nada — solo devuelve los campos encontrados para pre-llenar el formulario.
    """
    nombre = file.filename or ""
    ext = nombre.rsplit(".", 1)[-1].lower() if "." in nombre else ""
    if ext not in ("pdf", "xlsx", "xls", "docx"):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Formato no soportado. Use PDF, Excel (.xlsx) o Word (.docx).",
        )

    contenido = await file.read()
    texto = _extraer_texto(contenido, ext)
    resultado = _parsear_campos_rut(texto)
    resultado.nombre_archivo = nombre
    return resultado
