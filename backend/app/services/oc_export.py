"""Excel de solicitudes OC y filtro de búsqueda del listado.

Una hoja, tres consultas (solicitudes, cotizaciones, órdenes). Sin estilo por celda.
"""

import io
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import exists, or_
from sqlalchemy.orm import defer
from sqlmodel import Session, select

from app.models.oc import CotizacionProveedor, OrdenCompra, SolicitudOC

_BOG = ZoneInfo("America/Bogota")

_ESTADO = {
    "nueva": "Nueva",
    "en_cotizacion": "En cotización",
    "pendiente_aprobacion": "Cotización lista",
    "aprobada": "Aprobada",
    "rechazada": "Rechazada",
    "cancelada": "Cancelada",
    "en_correccion": "En corrección",
    "oc_enviada": "OC Enviada",
    "oc_en_plataforma": "En plataforma",
    "entregada": "Entregada",
    "cerrada": "Cerrada",
}

_HEADERS = [
    "Consecutivo OS",
    "Estado",
    "Prioridad",
    "Tipo",
    "Descripción",
    "Cantidad",
    "Categoría",
    "Grupo de artículos",
    "Solicitante",
    "Correo solicitante",
    "Área",
    "Sede",
    "Cliente",
    "Plataforma",
    "Placa / ficha",
    "Condición",
    "Proveedores",
    "NIT proveedor",
    "Correo proveedor",
    "Números de cotización",
    "Valor antes de IVA",
    "IVA",
    "Valor total",
    "Valor aprobado",
    "Fecha estimada entrega",
    "Número OC",
    "Número factura",
    "Número remisión",
    "Proforma",
    "Fecha solicitud",
    "Fecha asignación",
    "Fecha cotización",
    "Fecha aprobación",
    "Fecha envío OC",
    "Fecha en plataforma",
    "Fecha recibido",
    "Fecha cerrado",
    "Observaciones solicitante",
    "Observaciones compras",
    "Aval de compra",
]

_WIDTHS = [
    18, 18, 12, 14, 40, 10, 18, 22, 24, 28,
    18, 16, 22, 16, 16, 18, 28, 16, 28, 24,
    16, 14, 16, 16, 18, 16, 16, 16, 12, 18,
    18, 18, 18, 18, 18, 18, 18, 36, 36, 22,
]


def like_term(q: str) -> str:
    escaped = q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
    return f"%{escaped}%"


def cond_busqueda(q: str):
    """OS (consecutivo), nombre de proveedor o número de cotización. Sin distinguir mayúsculas."""
    term = like_term(q)
    cot = exists(
        select(CotizacionProveedor.id).where(
            CotizacionProveedor.solicitud_id == SolicitudOC.id,
            or_(
                CotizacionProveedor.proveedor_nombre.ilike(term, escape="\\"),
                CotizacionProveedor.numero_cotizacion_proveedor.ilike(term, escape="\\"),
            ),
        )
    )
    return or_(SolicitudOC.consecutivo_os.ilike(term, escape="\\"), cot)


def _ordenadas(cots: list) -> list:
    return sorted(cots, key=lambda c: c.created_at)


def texto_proveedores(cots: list) -> str:
    vistos: list[str] = []
    for c in _ordenadas(cots):
        nombre = (c.proveedor_nombre or "").strip()
        if nombre and nombre not in vistos:
            vistos.append(nombre)
    return ", ".join(vistos)


def texto_cotizaciones(cots: list) -> str:
    vistos: list[str] = []
    for c in _ordenadas(cots):
        num = (c.numero_cotizacion_proveedor or "").strip()
        if num and num not in vistos:
            vistos.append(num)
    return ", ".join(vistos)


def cotizacion_valores(cots: list):
    """Cotización aprobada; si no hay, la más reciente. De ahí salen NIT, valores y correo."""
    if not cots:
        return None
    aprobadas = [c for c in cots if c.aprobada is True]
    return max(aprobadas or cots, key=lambda c: c.created_at)


def _fmt_dt(value: datetime | None) -> str:
    if value is None:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(_BOG).strftime("%Y-%m-%d %H:%M")


def _fmt_fecha(value: date | None) -> str:
    return value.isoformat() if value else ""


def _num(value: float | None):
    return None if value is None else round(float(value), 2)


def _blank(value: str | None) -> str:
    return value or ""


def fila_excel(s: SolicitudOC, cots: list, orden: OrdenCompra | None) -> list:
    cot = cotizacion_valores(cots)
    entrega = s.fecha_estimada_entrega or (cot.fecha_estimada_entrega if cot else None)
    return [
        s.consecutivo_os,
        _ESTADO.get(s.estado, s.estado),
        s.nivel_prioridad,
        s.tipo_solicitud,
        s.descripcion,
        s.cantidad,
        _blank(s.categoria),
        _blank(s.grupo_articulos),
        s.solicitante_nombre,
        _blank(s.solicitante_email),
        _blank(s.area_solicitante),
        _blank(s.sede),
        _blank(s.cliente),
        _blank(s.plataforma),
        _blank(s.placa_ficha),
        _blank(s.condicion),
        texto_proveedores(cots),
        _blank(cot.proveedor_nit) if cot else "",
        _blank(cot.proveedor_email) if cot else "",
        texto_cotizaciones(cots),
        _num(cot.valor_antes_iva) if cot else None,
        _num(cot.valor_iva) if cot else None,
        _num(cot.valor_total) if cot else None,
        _num(cot.valor_aprobado) if cot else None,
        _fmt_fecha(entrega),
        orden.numero_oc if orden else "",
        _blank(s.numero_factura),
        _blank(s.numero_remision),
        "Sí" if s.tiene_proforma else "No",
        _fmt_dt(s.fecha_solicitud),
        _fmt_dt(s.fecha_asignacion),
        _fmt_dt(s.fecha_cotizacion),
        _fmt_dt(s.fecha_aprobacion),
        _fmt_dt(s.fecha_envio_oc),
        _fmt_dt(s.fecha_en_plataforma),
        _fmt_dt(s.fecha_recibido),
        _fmt_dt(s.fecha_cerrado),
        _blank(s.observaciones_solicitante),
        _blank(s.observaciones_compras),
        _blank(s.aval_compra),
    ]


def _indexar(rows: list, attr: str) -> dict:
    out: dict = {}
    for row in rows:
        key = getattr(row, attr)
        out.setdefault(key, []).append(row)
    return out


def workbook_bytes(filas: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Solicitudes OC"
    ws.append(_HEADERS)
    for fila in filas:
        ws.append(fila)
    for idx, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(_HEADERS))}{len(filas) + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar(oc_db: Session) -> tuple[bytes, str]:
    """Todas las solicitudes no archivadas. Tres lecturas, sin N+1."""
    sols = oc_db.exec(
        select(SolicitudOC)
        .where(SolicitudOC.archivada == False)  # noqa: E712
        .options(defer(SolicitudOC.fotos_producto))
        .order_by(SolicitudOC.fecha_solicitud.desc(), SolicitudOC.id.desc())
    ).all()
    ids = select(SolicitudOC.id).where(SolicitudOC.archivada == False)  # noqa: E712
    cots = oc_db.exec(
        select(CotizacionProveedor).where(CotizacionProveedor.solicitud_id.in_(ids))
    ).all()
    ordenes = oc_db.exec(
        select(OrdenCompra).where(OrdenCompra.solicitud_id.in_(ids))
    ).all()
    por_cot = _indexar(cots, "solicitud_id")
    por_oc: dict = {}
    for orden in ordenes:
        prev = por_oc.get(orden.solicitud_id)
        if prev is None or orden.created_at >= prev.created_at:
            por_oc[orden.solicitud_id] = orden
    filas = [fila_excel(s, por_cot.get(s.id, []), por_oc.get(s.id)) for s in sols]
    nombre = f"solicitudes-oc-{datetime.now(_BOG).strftime('%Y%m%d')}.xlsx"
    return workbook_bytes(filas), nombre


def _selfcheck() -> None:
    from types import SimpleNamespace

    assert like_term("COT-150") == "%COT-150%"
    assert like_term("100%") == "%100\\%%"
    assert like_term("a_b") == "%a\\_b%"
    t0 = datetime(2026, 1, 1, tzinfo=timezone.utc)
    t1 = datetime(2026, 2, 1, tzinfo=timezone.utc)
    aprobada = SimpleNamespace(
        aprobada=True, created_at=t0, proveedor_nombre="Pepito",
        numero_cotizacion_proveedor="COT-1", proveedor_nit="900", proveedor_email="a@b.c",
        valor_antes_iva=8, valor_iva=2, valor_total=10, valor_aprobado=9,
        fecha_estimada_entrega=None,
    )
    otra = SimpleNamespace(
        aprobada=None, created_at=t1, proveedor_nombre="Otro",
        numero_cotizacion_proveedor="COT-150", proveedor_nit=None, proveedor_email=None,
        valor_antes_iva=None, valor_iva=None, valor_total=99, valor_aprobado=None,
        fecha_estimada_entrega=None,
    )
    assert cotizacion_valores([otra, aprobada]).proveedor_nombre == "Pepito"
    assert texto_proveedores([otra, aprobada]) == "Pepito, Otro"
    assert texto_cotizaciones([otra, aprobada]) == "COT-1, COT-150"
    sola = SimpleNamespace(
        aprobada=None, created_at=t1, proveedor_nombre="Pepito",
        numero_cotizacion_proveedor="COT-150", proveedor_nit=None, proveedor_email=None,
        valor_antes_iva=None, valor_iva=None, valor_total=5, valor_aprobado=None,
        fecha_estimada_entrega=None,
    )
    assert cotizacion_valores([sola]).numero_cotizacion_proveedor == "COT-150"
    s = SimpleNamespace(
        consecutivo_os="OS-2026-0001", estado="nueva", nivel_prioridad="Alta",
        tipo_solicitud="compra", descripcion="Tornillos", cantidad=2,
        categoria="Repuestos", grupo_articulos="Ferretería", solicitante_nombre="Ana",
        solicitante_email=None, area_solicitante=None, sede=None, cliente=None,
        plataforma=None, placa_ficha=None, condicion=None, fecha_estimada_entrega=None,
        numero_factura=None, numero_remision=None, tiene_proforma=False,
        fecha_solicitud=t0, fecha_asignacion=None, fecha_cotizacion=None,
        fecha_aprobacion=None, fecha_envio_oc=None, fecha_en_plataforma=None,
        fecha_recibido=None, fecha_cerrado=None, observaciones_solicitante=None,
        observaciones_compras=None, aval_compra=None,
    )
    fila = fila_excel(s, [aprobada, otra], None)
    assert len(fila) == len(_HEADERS) == len(_WIDTHS)
    assert fila[0] == "OS-2026-0001"
    assert fila[16] == "Pepito, Otro"
    assert fila[19] == "COT-1, COT-150"
    assert fila[23] == 9
    blob = workbook_bytes([fila])
    assert blob[:2] == b"PK"


if __name__ == "__main__":
    _selfcheck()
